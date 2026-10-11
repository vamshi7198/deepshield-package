"""Builds the improved IP Pre-DOI patent deck from the inventors' template.

Keeps the template's title, inventor and closing slides (unchanged) and its look (Cambria, black titles,
bordered tables); rebuilds the content slides with accurate claims, native editable diagrams (novel blocks
in red) and the verified evaluation evidence.
Usage: tools\\.venv\\Scripts\\python tools\\build_patent_deck.py
"""
import copy
import json
import os

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(ROOT, "DeepSheild_Templete_IP_PreDOI Doc Updated Revised V0.3 09.2026 (1).pptx")
OUT = os.path.join(ROOT, "deliverables", "DeepShield_Patent_PreDOI_v1.0.pptx")
SCREENS = os.path.join(ROOT, "deliverables", "_work", "screens")
FIGS = os.path.join(ROOT, "deliverables", "_work", "figs")
R2 = os.path.join(ROOT, "eval", "runs", "round2", "deepshield_round2")

FONT = "Cambria"
BLACK = RGBColor(0, 0, 0)
RED = RGBColor(0xC0, 0x00, 0x00)
RED_FILL = RGBColor(0xFD, 0xEC, 0xEC)
GREY = RGBColor(0x59, 0x59, 0x59)
LIGHT = RGBColor(0xF2, 0xF2, 0xF2)
BLUE_FILL = RGBColor(0xE8, 0xEE, 0xF7)
BLUE_LINE = RGBColor(0x2F, 0x52, 0x8F)
GREEN = RGBColor(0x0E, 0x7A, 0x4F)
AMBER = RGBColor(0xB2, 0x6B, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

summary = json.load(open(os.path.join(R2, "summary.json"), encoding="utf-8"))
matched = json.load(open(os.path.join(ROOT, "deliverables", "_work", "matched_review_comparison.json")))
combined = json.load(open(os.path.join(ROOT, "deliverables", "_work", "combined_review_score.json")))

prs = Presentation(TEMPLATE)
LAYOUT = prs.slide_layouts[1]  # 제목 및 내용 (title and content), as used by every template slide

# ----------------------------------------------------------------------------- keep 1, 2, 22; drop the rest
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
keep = {0: "title", 1: "inventors", 21: "thanks"}
kept = {}
for i, sid in enumerate(ids):
    if i in keep:
        kept[keep[i]] = sid
    else:
        prs.part.drop_rel(sid.get(qn("r:id")))
        sldIdLst.remove(sid)
# give the kept slides (and their notes) names that new slides cannot collide with
from pptx.opc.packuri import PackURI
for n, sid in enumerate(list(sldIdLst), start=901):
    part = prs.part.related_part(sid.get(qn("r:id")))
    part.partname = PackURI(f"/ppt/slides/slide{n}.xml")
    if part.has_notes_slide:
        part.notes_slide.part.partname = PackURI(f"/ppt/notesSlides/notesSlide{n}.xml")


# ----------------------------------------------------------------------------- helpers
def new_slide(title):
    s = prs.slides.add_slide(LAYOUT)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    if title:
        tb = s.shapes.add_textbox(Inches(0.3), Inches(0.08), Inches(12.7), Inches(0.75))
        tb.name = "Title"
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        p._p.get_or_add_pPr().set("latinLnBrk", "0")
        r = p.add_run()
        r._r.get_or_add_rPr().set("lang", "en-US")
        r.text = title
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(28), True, BLACK
    return s


def runs_into(p, parts, size, color=BLACK, bold=False):
    """parts: str or list of (text, style) where style in {'', 'b', 'r', 'rb', 'i', 'g'}."""
    if isinstance(parts, str):
        parts = [(parts, "")]
    p._p.get_or_add_pPr().set("latinLnBrk", "0")   # template is Korean-default: never split Latin words
    for txt, st in parts:
        r = p.add_run()
        r._r.get_or_add_rPr().set("lang", "en-US")
        r.text = txt
        f = r.font
        f.name, f.size = FONT, Pt(size)
        f.bold = bold or ("b" in st)
        f.italic = "i" in st
        f.color.rgb = RED if "r" in st else (GREY if "g" in st else color)


def textbox(s, x, y, w, h, paras, size=14, color=BLACK, bullet=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=4):
    """paras: list of paragraphs; each paragraph is str or list of (text, style); ('##', text) = bold heading line."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, Inches(0.04))
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(space)
        if isinstance(para, tuple) and para[0] == "##":
            runs_into(p, para[1], size + 1, color, bold=True)
            continue
        is_bullet = bullet
        if isinstance(para, tuple) and para[0] == "-":
            is_bullet, para = True, para[1]
        if is_bullet:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Inches(0.22)))
            pPr.set("indent", str(-Inches(0.2)))
            bu = etree.SubElement(pPr, qn("a:buChar"))
            bu.set("char", "•")
        runs_into(p, para, size, color)
    return tb


def box(s, x, y, w, h, fill=WHITE, line=BLACK, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, dash=False):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = 0.12
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
        if dash:
            sh.line.dash_style = 4  # dash
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh


def label_box(s, x, y, w, h, title, sub=None, novel=False, proposed=False, size=12):
    fill, line = (RED_FILL, RED) if novel else (BLUE_FILL, BLUE_LINE)
    if proposed:
        fill, line = WHITE, GREY
    b = box(s, x, y, w, h, fill=fill, line=line, lw=1.75 if novel else 1.0, dash=proposed)
    tf = b.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ("margin_left", "margin_right"):
        setattr(tf, m, Inches(0.06))
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    runs_into(p, title, size, RED if novel else (GREY if proposed else BLACK), bold=True)
    if sub:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        runs_into(p2, sub, size - 2, GREY)
    return b


def arrow(s, x1, y1, x2, y2, color=GREY, w=1.25):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    return c


def set_cell_border(cell, color="000000", width_pt=0.75):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for old in tcPr.findall(qn(tag)):
            tcPr.remove(old)
        ln = etree.SubElement(tcPr, qn(tag), w=str(int(width_pt * 12700)), cap="flat", cmpd="sng", algn="ctr")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr"), val=color)
        etree.SubElement(ln, qn("a:prstDash"), val="solid")


def table(s, x, y, w, col_w, rows, size=12, header_fill=LIGHT, row_h=None, first_col_bold=False):
    """rows[0] is the header; every cell is str or list of (text, style)."""
    nr, nc = len(rows), len(rows[0])
    gt = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(0.3 * nr))
    tbl = gt.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("bandRow", "0")
    tblPr.set("firstRow", "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # 'No Style, Table Grid'
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        if row_h:
            tbl.rows[i].height = Inches(row_h if i else min(0.5, max(0.32, row_h * 0.7)))
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = header_fill if i == 0 else WHITE
            set_cell_border(c)
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = ""
            runs_into(p, val, size, BLACK, bold=(i == 0) or (first_col_bold and j == 0))
    return tbl


def picture(s, path, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return s.shapes.add_picture(path, Inches(x), Inches(y), **kw)


def legend(s, x=10.2, y=0.3):
    b = box(s, x, y + 0.04, 0.22, 0.22, fill=RED_FILL, line=RED, lw=1.5, shape=MSO_SHAPE.OVAL)
    textbox(s, x + 0.3, y - 0.02, 1.0, 0.35, [[("Novel block", "r")]], size=12)
    box(s, x + 1.45, y + 0.04, 0.22, 0.22, fill=BLUE_FILL, line=BLUE_LINE, shape=MSO_SHAPE.OVAL)
    textbox(s, x + 1.75, y - 0.02, 1.2, 0.35, [[("Established", "")]], size=12)


def pct(v):
    return f"{100 * v:.1f}%"


slides = {}  # name -> slide (for the table of contents)
M = summary["main_test"]
ITW = summary["in_the_wild"]
LOGO = {r["held_out"]: r for r in summary["logo"]}

# ============================================================================= 3. TOC (filled at the end)
toc = new_slide("Table of Content")
slides["toc"] = toc

# ============================================================================= 4. Recommendation summary
s = new_slide("Recommendation Summary")
slides["Recommendation Summary"] = s
tbl = table(s, 0.3, 0.9, 12.7, [3.4, 4.3, 5.0], [["Summary of the Invention", "Representative Drawing", "Evaluation Opinion"], ["", "", ""]], size=14)
tbl.rows[1].height = Inches(6.15)
for j in range(3):
    tbl.cell(0, j).text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
textbox(s, 0.35, 1.45, 3.3, 5.5, [
    "An audio deepfake screening system that goes beyond a binary real/fake call.",
    "A two-channel (Mel-spectrogram + MFCC) picture of each 3-second window is read by a compact residual CNN with Squeeze-and-Excitation attention, giving one shared 128-number embedding.",
    [("From that one embedding, four heads run in parallel: ", ""), ("spoof score, attack-type, MC-dropout uncertainty and open-set novelty", "r"), (".", "")],
    [("A ", ""), ("risk-aware gateway", "r"), (" combines them to route each clip to Pass, Review, Block or Unknown-pattern review, with a plain-language rationale.", "")],
], size=12, space=6)
# representative drawing: vertical pipeline
steps = [
    ("1  Audio input", "any format → 16 kHz, 3-s windows", False, False),
    ("2  Robustness augmentation", "training: MP3 · noise · reverb · 8 kHz phone", False, False),
    ("3  Dual-channel features", "Mel(64) + MFCC(64) → 2×64×94", False, False),
    ("4  Residual CNN + SE attention", "3 residual blocks · 341K parameters", False, False),
    ("5  Shared 128-d embedding", "one representation feeds all heads", False, False),
    ("6  Four parallel heads", "spoof · attack-type · uncertainty · novelty", True, False),
    ("7  Risk-aware gateway", "Pass / Review / Block / Unknown", True, False),
    ("8  Explanation layer", "deterministic rationale (+ per-channel CAM*)", True, False),
    ("9  Continual-learning loop", "proposed, not yet built", False, True),
]
bx, bw, bh, by0, gap = 4.05, 3.9, 0.5, 1.42, 0.14
for k, (t, sub, nov, prop) in enumerate(steps):
    yy = by0 + k * (bh + gap)
    label_box(s, bx, yy, bw, bh, t, sub, novel=nov, proposed=prop, size=11)
    if k < len(steps) - 1:
        arrow(s, bx + bw / 2, yy + bh, bx + bw / 2, yy + bh + gap)
textbox(s, 0.35, 6.25, 3.3, 0.75, [[("Drawing: red = novel blocks · dashed = proposed · *code exists, not evaluated", "g")]], size=10)
textbox(s, 8.35, 1.45, 4.55, 5.6, [
    ("##", [("Prior-art opinion", "")]),
    "Each component (Mel/MFCC features, ResNet + SE, MC-dropout uncertainty, Mahalanobis out-of-distribution scoring, reject-option review) is individually known.",
    ("##", [("What is new", "")]),
    [("The ", ""), ("combination", "r"), (": four heads computed from one shared embedding and ", ""), ("routed jointly", "r"), (" by a gateway calibrated to a human-review budget.", "")],
    ("##", [("Evidence so far (verified)", "")]),
    [("A review lane at ~15% genuine workload halves both error types on never-seen generators. The heads' ", ""), ("extra", "b"), (" value over a score-only review rule is small but consistent (9 of 9 unseen families).", "")],
    ("##", [("Assessment", "")]),
    [("Inventors proposed Class A1. ", ""), ("Current evidence supports a narrower combination claim; the inventive step is at risk", "r"), (" until the heads' contribution is shown more strongly (slides 22–23, 26).", "")],
], size=12, space=3)

# ============================================================================= 5. Significance
s = new_slide("Opportunity Review: Significance")
slides["Opportunity Review"] = s
textbox(s, 0.35, 0.95, 7.4, 6.3, [
    ("##", "What is the idea?"),
    "Instead of asking only “is this audio fake?”, the system answers five questions:",
    ("-", "Is it fake?  ·  What kind of fake?  ·  How confident are we?"),
    ("-", "Have we seen anything like this before?  ·  Why did we decide this?"),
    ("##", "Why is it useful?"),
    "Voice cloning improves every few months. A detector trained on today’s generators can fail on tomorrow’s, so a deployable system must recognise when it is unsure and hand the case to a person, and give a defensible reason for every block.",
    ("##", "Significance"),
    ("-", "Routes doubtful audio to human review instead of silently calling it “real”"),
    ("-", "A reviewer sees the decision, scores, attack family, the most suspicious 3 s and a one-line reason"),
    ("-", "Tested against MP3 compression, additive noise, reverberation and 8 kHz telephone audio"),
    ("-", "Runs on a laptop CPU (~341K parameters) in well under a second per clip"),
    ("-", "Pain point addressed: voice-fraud calls and contested audio evidence"),
], size=14, space=5)
picture(s, os.path.join(SCREENS, "result_review_uncertain_card.png"), 8.0, 1.0, h=5.75)
textbox(s, 8.0, 6.8, 4.9, 0.4, [[("Live system: an AI-generated clip routed to human review because uncertainty is high", "g")]], size=10)

# ============================================================================= 6. Current solutions + pain points
s = new_slide("Opportunity Review: Current Solutions & Pain Points")
textbox(s, 0.35, 0.95, 6.2, 6.2, [
    ("##", "What are the current solutions?"),
    ("-", "Single-channel spectrogram CNN detectors"),
    ("-", "SE-ResNet / LCNN classifiers adapted from image recognition"),
    ("-", "Static single-threshold binary classifiers (one probability, one cut-off)"),
    ("-", "Closed-set classifiers that assume a fixed list of generators"),
    ("-", "Self-supervised front-ends (wav2vec 2.0 / WavLM) with a binary head"),
    ("-", "Generic Grad-CAM / SHAP saliency for explanation"),
], size=14, space=5)
textbox(s, 6.75, 0.95, 6.25, 6.2, [
    ("##", "Pain points of current solutions"),
    ("-", "Output is only real/fake: no attack family, no confidence, no reason"),
    ("-", "Closed-set training: unseen generators are often silently classified as real"),
    ("-", [("Accuracy on random splits is inflated. DeepShield’s own earlier 97.7% fell to ", ""), ("93.2%", "b"), (" once recordings, sentences and speakers were kept apart (28% of random-split validation clips had a copy in training)", "")]),
    ("-", "Compression, noise, reverberation and device variation are rarely tested"),
    ("-", "Generic saliency maps do not say why audio is synthetic"),
    ("##", "Key gaps"),
    ("-", "No joint attribution, novelty and uncertainty routing"),
    ("-", "No human-review lane with a controlled workload"),
    ("-", "No defensible, per-decision explanation"),
], size=14, space=5)

# ============================================================================= 7. Proposed engine (with build status)
s = new_slide("Opportunity Review: Proposed Engine and Build Status")
rows = [["Module", "What it does", "Status"]]
st_rows = [
    ("Robustness-aware augmentation", "MP3 32–128 kbps, noise 0–20 dB SNR, synthetic reverb, 8 kHz telephone band", "Implemented & evaluated"),
    ("Dual-channel feature grid", "Mel(64) + MFCC(64) stacked → 2 × 64 × 94 per 3-s window", "Implemented & evaluated"),
    ("Residual CNN + SE attention", "3 residual blocks, 32→64→128 channels, SE after block 2; 341,254 parameters", "Implemented & evaluated"),
    ("Shared 128-d embedding", "Global-pooled output read by every head", "Implemented & evaluated"),
    ([("Four parallel heads", "r")], "Spoof score · attack-type (5 classes) · MC-dropout uncertainty · Mahalanobis novelty", "Implemented & evaluated"),
    ([("Risk-aware gateway", "r")], "Novelty → uncertainty → probability bands; uncertainty threshold calibrated to a review budget", "Implemented & evaluated"),
    ([("Explanation layer", "r")], "Deterministic rationale (live). Per-channel Mel-vs-MFCC CAM disagreement map", "Rationale live; CAM code not evaluated"),
    ("Continual-learning loop", "Reviewer-confirmed + high-novelty samples → periodic fine-tune", "Proposed, not built"),
]
for a, b, c in st_rows:
    rows.append([a, b, [(c, "b" if c.startswith("Implemented &") else "r")]])
table(s, 0.3, 0.95, 12.7, [3.2, 6.6, 2.9], rows, size=13, row_h=0.6)
textbox(s, 0.35, 6.5, 12.6, 0.8, [[("Advantages: ", "b"), ("answers “what kind of fake” as well as “is it fake”; flags unfamiliar inputs; confidence from MC-dropout; review workload is a tunable parameter; every decision carries a deterministic reason.", "")]], size=13)

# ============================================================================= 8. Why it works
s = new_slide("Opportunity Review: Why the New Engine Works")
textbox(s, 0.35, 0.95, 6.6, 6.3, [
    [("Dual-channel fusion. ", "b"), ("The Mel channel exposes fine spectral artefacts left by vocoders; the MFCC channel summarises the vocal-tract envelope. Stacking them lets one network reason across both.", "")],
    [("One shared embedding. ", "b"), ("All heads judge the same 128-d representation, so their outputs are consistent and inference stays cheap.", "")],
    [("Uncertainty by repeated stochastic passes. ", "b"), ("Dropout is kept active at inference and the spoof head is queried N times; the spread of answers is the uncertainty. Wrong verdicts average 0.37 vs 0.17 for correct ones.", "")],
    [("Genuine-only novelty boundary. ", "b"), ("A Gaussian is fitted to embeddings of authentic speech only. On seen generators novelty carries no signal (AUC 0.55), but on never-seen generators it does (AUC 0.60–0.78), which is the intended open-set effect.", "")],
    [("Review lane with a fixed budget. ", "b"), ("Doubtful clips go to people; the threshold is set so ≤15% of genuine validation audio is reviewed.", "")],
    [("Removed from this version: ", "r"), ("the earlier claim about “invariant cues” (breath, jitter/shimmer, glottal periodicity) is not implemented and is moved to future work.", "")],
], size=13, space=8)
picture(s, os.path.join(FIGS, "two_views_genuine_vs_fake.png"), 7.15, 1.05, w=5.9)
textbox(s, 7.15, 4.2, 5.9, 0.6, [[("The two channels for a genuine clip (left) and an AI-generated clip (right). Note the unnaturally clean silences in the fake.", "g")]], size=10)

# ============================================================================= 9. Recommendation
s = new_slide("Recommendation")
quad = [
    ("1) What is the core technology in this idea?", "A dual-channel (Mel + MFCC) residual CNN with SE attention producing one shared embedding, read in parallel by four heads (spoof, attack-type, uncertainty, novelty), with a risk-aware gateway and explanation layer on top."),
    ("2) What complex problem is solved?", "Screening speech audio when new voice generators keep appearing: the system must still catch most fakes, avoid accusing genuine speakers, and say “I am not sure” in a controlled number of cases."),
    ("3) Why could it be game-changing?", "It turns a one-number classifier into a trust workflow: detect, attribute, estimate confidence, flag the unfamiliar, explain, and route to people with a predictable workload."),
    ("4) What is the breakthrough, and how strong is the evidence?", [("The jointly routed gateway. Verified: a ~15% review lane halves both error types on unseen generators. ", ""), ("Not yet shown strongly: that the uncertainty/novelty heads beat a score-only review rule at the same workload.", "r")]),
]
for k, (h, body) in enumerate(quad):
    x = 0.3 + (k % 2) * 6.45
    y = 0.95 + (k // 2) * 3.15
    qt = table(s, x, y, 6.25, [6.25], [[h], [body]], size=13)
    qt.rows[0].height, qt.rows[1].height = Inches(0.5), Inches(2.4)

# ============================================================================= 10. Independent claims
s = new_slide("Independent Claims")
slides["Independent Claims"] = s
textbox(s, 7.6, 0.18, 5.5, 0.4, [[("In order of relevancy. Novel features highlighted in ", "g"), ("red", "r"), (".", "g")]], size=11)
textbox(s, 0.35, 0.95, 5.0, 6.3, [
    ("##", "Abstract"),
    "A speech-audio screening method and system that detects synthetic speech, attributes the generator family, estimates its own uncertainty, flags inputs unlike authentic speech, and routes each input to automatic pass, automatic block or human review, with a deterministic explanation, all from one shared dual-channel embedding.",
    ("##", "Use cases"),
    ("-", "Voice-fraud screening of calls and voice notes"),
    ("-", "Call-centre and voice-biometric authentication"),
    ("-", "Forensic triage of recorded evidence"),
    ("-", "Compliance workflows with a human review queue"),
], size=13, space=5)
s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(5.55), Inches(0.95), Inches(5.55), Inches(7.2)).line.color.rgb = GREY
textbox(s, 5.75, 0.95, 7.3, 6.4, [
    ("##", "Claim 1 (method)"),
    "A computer-implemented method for screening speech audio, comprising:",
    "(a) dividing the audio into fixed-length windows and forming, for each window, a two-channel time–frequency array whose first channel is a log-power Mel spectrogram and whose second channel is a matrix of Mel-frequency cepstral coefficients;",
    "(b) processing the array with a residual convolutional network having channel attention to obtain a single shared embedding;",
    [("(c) computing from that same shared embedding, in parallel, a spoof probability, an attack-family attribution, ", ""), ("a predictive-uncertainty estimate obtained from repeated stochastic forward passes, and a novelty score measuring the distance of the embedding from a distribution fitted only on embeddings of authentic speech", "r"), (";", "")],
    [("(d) ", ""), ("routing the audio to one of automatic pass, automatic block, human review or unknown-pattern review as a joint function of the spoof probability, the uncertainty estimate and the novelty score, wherein the novelty and uncertainty conditions take precedence over the spoof-probability thresholds", "r"), ("; and", "")],
    "(e) outputting the routing decision together with an explanation generated from the head outputs.",
    ("##", "Claim 2 (system)"),
    [("A system comprising one or more processors and memory storing instructions that implement a feature module, the residual network, the four heads and ", ""), ("a gateway module configured to perform step (d) of claim 1", "r"), (".", "")],
], size=12, space=4)

# ============================================================================= 11. Dependent claims
s = new_slide("Dependent Claims")
textbox(s, 7.6, 0.18, 5.5, 0.4, [[("Status in brackets: [E] implemented & evaluated · [C] code only · [P] proposed", "g")]], size=11)
deps = [
    [("Claim 3: ", "b"), ("The method of claim 1, wherein the novelty score is a Mahalanobis distance from a Gaussian fitted on authentic-speech embeddings, with the threshold at a percentile (97.5th) of authentic training distances, ", ""), ("checked before any probability threshold", "r"), (". [E]", "")],
    [("Claim 4: ", "b"), ("The method of claim 1, wherein the uncertainty estimate is the standard deviation of the spoof probability over N stochastic passes with dropout applied to the shared embedding and kept active at inference. [E]", "")],
    [("Claim 5: ", "b"), ("The method of claim 1, wherein ", ""), ("the uncertainty threshold is calibrated on held-out validation data so that at most a target fraction (e.g. 15%) of authentic inputs is routed to human review", "r"), (". [E]", "")],
    [("Claim 6: ", "b"), ("The method of claim 1, wherein ", ""), ("the routing uses a combined risk score: a weighted combination of the spoof-probability margin, the uncertainty estimate and the novelty score, with the weights fitted on validation data", "r"), (". [C: preliminary analysis, slide 23]", "")],
    [("Claim 7: ", "b"), ("The method of claim 1, wherein a long recording is scanned in several windows, the decision is based on the most suspicious window, and its time offset is reported to the reviewer. [E]", "")],
    [("Claim 8: ", "b"), ("The method of claim 1, wherein ", ""), ("audio shorter than a minimum duration (e.g. 2 s) is routed to human review", "r"), (" instead of receiving an automatic decision. [E: demo clips]", "")],
    [("Claim 9: ", "b"), ("The method of claim 1, wherein the attack-family head classifies among real, vocoder artefact, TTS prosody, splice/edit and unknown-synthetic, trained jointly with the spoof head. [E, splice/edit class untrained]", "")],
    [("Claim 10: ", "b"), ("The method of claim 1, wherein the explanation is a deterministic template assembled from the head outputs, never generated free-form. [E]", "")],
    [("Claim 11: ", "b"), ("The method of claim 1, further comprising ", ""), ("per-channel activation maps computed separately for the Mel and MFCC channels, whose disagreement map is shown to the reviewer", "r"), (". [C]", "")],
    [("Claim 12: ", "b"), ("The method of claim 1, further comprising training with randomised MP3 re-encoding, additive noise at controlled SNR, reverberation and 8 kHz band-limiting. [E]", "")],
    [("Claim 13: ", "b"), ("The method of claim 1, further comprising queuing reviewer-confirmed and high-novelty samples for periodic fine-tuning with versioned model releases. [P]", "")],
]
textbox(s, 0.35, 0.85, 12.65, 6.5, deps, size=12, space=5)

# ============================================================================= 12. Main patentable points (ranked, with evidence)
s = new_slide("Main Patentable Points (ranked by evidence)")
slides["Main Patentable Points"] = s
table(s, 0.3, 0.95, 12.7, [0.4, 4.6, 5.6, 2.1], [
    ["#", "Patentable point", "Evidence from our evaluation", "Strength"],
    ["1", [("Joint routing from one embedding using spoof score, uncertainty and novelty", "r")], "Review lane at ~15% workload halves errors on unseen generators. Over a score-only review rule at equal workload: current rules no better; a fitted combination is better in 9/9 families but by a small margin (error-spotting AUC 0.79 → 0.81).", [("Moderate–weak", "rb")]],
    ["2", [("Workload-calibrated review threshold", "r")], "Calibrated on validation to ≤15%; achieved 13.8% on test, 14.4% on unseen generators. Calibration as such is a known technique.", [("Moderate", "b")]],
    ["3", [("Genuine-only novelty boundary checked first", "r")], "No signal on seen generators (AUC 0.55); real signal on unseen ones (0.60–0.78). Flags 0–31% of unseen fakes at the current threshold.", [("Moderate–weak", "rb")]],
    ["4", "Attack-family attribution trained jointly with detection", "91.0% correct on test fakes; classes are coarse and the splice/edit class has no training data.", [("Moderate", "b")]],
    ["5", [("Per-channel (Mel vs MFCC) disagreement explanation", "r")], "Code exists; never evaluated or shown to users.", [("None yet", "rb")]],
], size=12, row_h=0.9)
textbox(s, 0.35, 6.75, 12.6, 0.6, [[("Recommendation to the patent team: ", "b"), ("lead with point 1 as a combination claim and support it with point 2. Points 3 and 5 need the experiments on slide 26 before they can carry weight.", "")]], size=13)

# ============================================================================= 13. Novelty map
s = new_slide("Novelty Map: what is new, and why")
slides["Novelty Map"] = s
table(s, 0.3, 0.95, 12.7, [3.1, 4.2, 1.6, 3.8], [
    ["Block", "Closest known technique", "Novel here?", "Why / why not"],
    ["Mel + MFCC two-channel input", "Common hand-crafted front-ends in ASVspoof systems", "No", "Established; chosen for complementary cues"],
    ["Residual CNN + SE attention", "ResNet / SE-Net (image recognition), widely used in anti-spoofing", "No", "Established backbone, deliberately compact"],
    ["MC-dropout uncertainty", "Gal & Ghahramani, 2016", "No (alone)", "Known estimator; new in how it drives routing"],
    ["Mahalanobis novelty score", "Lee et al., 2018 (out-of-distribution detection)", "No (alone)", "Known; new in being fitted on authentic speech only and checked first in routing"],
    ["Human-review lane / reject option", "Selective classification (Geifman & El-Yaniv, 2017)", "No (alone)", "Known; DeepShield sets it by a review budget"],
    [[("Four heads on one shared embedding + jointly routed gateway", "r")], "No single reference combining all four signals in one routing function (per inventors’ search)", [("Candidate", "rb")], "Integrated combination; effect modest so far (slides 22–23)"],
    [[("Per-channel disagreement explanation", "r")], "Grad-CAM (Selvaraju et al., 2017) applied per input channel", [("Candidate", "rb")], "Specific to the two-channel design; not yet evaluated"],
], size=12, row_h=0.62)
textbox(s, 0.35, 6.7, 12.6, 0.6, [[("Academic references are well-known papers named for orientation; a formal prior-art search is still required.", "g")]], size=11)

# ============================================================================= 14. Prior art comparison
s = new_slide("Prior Art Comparison")
slides["Prior Art Comparison"] = s
table(s, 0.3, 0.95, 12.7, [5.0, 7.7], [
    ["Prior-art reference", "Difference from DeepShield"],
    ["Single-channel spectrogram CNN detectors (US 12,462,813 B1, cited by the inventors)*", "DeepShield fuses Mel + MFCC into one 2 × 64 × 94 input and adds attribution, uncertainty and novelty heads"],
    ["Static single-threshold binary classifiers (US 12,592,239 B2, cited by the inventors)*", "DeepShield routes on spoof score, uncertainty and novelty together, with a budgeted human-review lane"],
    ["Closed-set generator classifiers (ASVspoof 2019 LA challenge systems)", "DeepShield adds a novelty score fitted on authentic speech only and an Unknown-pattern route"],
    ["Selective classification / reject option (Geifman & El-Yaniv, 2017)", "DeepShield’s review decision uses head signals beyond the classifier’s own confidence"],
    ["Generic Grad-CAM / SHAP saliency (Selvaraju et al., 2017; Lundberg & Lee, 2017)", "DeepShield adds a deterministic rationale and proposes a Mel-vs-MFCC disagreement map"],
], size=13, row_h=0.9)
textbox(s, 0.35, 6.55, 12.6, 0.7, [[("* Patent numbers as cited in the inventors’ draft; not independently verified. To be confirmed by the patent team’s search.", "r")]], size=12)

# ============================================================================= 15. Difference from closest prior art
s = new_slide("Difference from Closest Prior Art")
table(s, 0.3, 0.95, 12.7, [2.5, 5.4, 4.8], [
    ["Feature", "Proposed (DeepShield)", "Prior art"],
    ["Output", "Spoof score + attack family + uncertainty + novelty from one shared embedding", "Single real/fake probability"],
    ["Unseen generators", "Novelty and uncertainty can divert doubtful inputs to review", "Closed set; often silently classified as real"],
    ["Decision routing", "Joint function of three signals; novelty checked first; review budget calibrated", "Fixed threshold on one score"],
    ["Workload control", "Genuine review rate set in advance (≤15% on validation)", "Not addressed"],
    ["Explainability", "Deterministic rationale; per-channel disagreement map (proposed)", "Generic saliency, or none"],
], size=13, row_h=0.72)
textbox(s, 0.35, 5.6, 12.6, 1.6, [
    ("##", "Summary of deficiencies in prior approaches"),
    "A bare real/fake score fails silently on unseen generators and degraded audio, gives no usable confidence, no workload control and no specific explanation. DeepShield addresses each of these structurally; the size of the gain attributable to the extra heads is still being established.",
], size=13)

# ============================================================================= 16. Section header
s = new_slide("")
textbox(s, 0.37, 3.3, 12.6, 0.9, [[("System Architecture & Module Description", "b")]], size=32)
slides["System Architecture"] = s

# ============================================================================= 17. End-to-end pipeline
s = new_slide("6. Systems and Methods: End-to-End Pipeline")
legend(s)
row1 = [("Audio input", "any format → 16 kHz"), ("Windowing", "3-s windows, up to 12 per clip"), ("Augmentation", "training only"), ("Dual-channel grid", "Mel + MFCC, 2×64×94"), ("Residual CNN + SE", "341K parameters"), ("Shared embedding", "128-d")]
xw, xg, y1 = 1.85, 0.27, 1.25
for k, (t, sub) in enumerate(row1):
    x = 0.35 + k * (xw + xg)
    label_box(s, x, y1, xw, 0.95, t, sub, size=12)
    if k < len(row1) - 1:
        arrow(s, x + xw, y1 + 0.47, x + xw + xg, y1 + 0.47)
heads = [("Spoof score", "P(fake)"), ("Attack-type", "5 families"), ("Uncertainty", "MC-dropout, N passes"), ("Novelty", "Mahalanobis, genuine-only")]
hx0, hy, hw, hg = 0.9, 2.95, 2.6, 0.35
emb_cx = 0.35 + 5 * (xw + xg) + xw / 2
arrow(s, emb_cx, y1 + 0.95, emb_cx, 2.55)
s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(hx0 + hw / 2), Inches(2.55), Inches(emb_cx), Inches(2.55)).line.color.rgb = GREY
for k, (t, sub) in enumerate(heads):
    x = hx0 + k * (hw + hg)
    arrow(s, x + hw / 2, 2.55, x + hw / 2, hy)
    label_box(s, x, hy, hw, 0.85, t, sub, novel=True, size=12)
gw = label_box(s, 3.15, 4.45, 5.6, 0.95, "Risk-aware gateway", "joint routing · novelty → uncertainty → bands · review budget", novel=True, size=13)
for k in range(4):
    x = hx0 + k * (hw + hg) + hw / 2
    arrow(s, x, hy + 0.85, 3.15 + 5.6 * (k + 0.5) / 4, 4.45)
outs = [("Pass", GREEN), ("Human review", AMBER), ("Unknown pattern", AMBER), ("Block", RED)]
for k, (t, col) in enumerate(outs):
    x = 0.9 + k * 2.45
    b = box(s, x, 5.95, 2.2, 0.6, fill=WHITE, line=col, lw=1.75)
    tf = b.text_frame
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    runs_into(tf.paragraphs[0], t, 13, col, bold=True)
    arrow(s, 3.15 + 5.6 * (k + 0.5) / 4, 5.4, x + 1.1, 5.95)
label_box(s, 9.45, 4.45, 3.5, 0.95, "Explanation layer", "rationale (+ per-channel CAM*)", novel=True, size=12)
arrow(s, 8.75, 4.92, 9.45, 4.92)
label_box(s, 10.75, 5.95, 2.2, 0.6, "Continual learning", "proposed", proposed=True, size=11)
arrow(s, 11.85, 5.4, 11.85, 5.95)
textbox(s, 0.35, 6.75, 12.6, 0.55, [[("Each window is scored; the most suspicious window drives the decision and its offset is shown to the reviewer. *Code exists, not evaluated.", "g")]], size=11)

# ============================================================================= 18. Gateway decision logic
s = new_slide("6. Systems and Methods: Gateway Decision Logic")
legend(s)
flow = [
    ("Audio shorter than 2 s?", "Review: too short to judge"),
    ("Novelty score > threshold (97.5th pct. of genuine)?", "Unknown pattern → human review"),
    ("Uncertainty > calibrated threshold (0.445)?", "High uncertainty → human review"),
    ("Spoof probability < 0.30?", "Pass (automatic clearance)"),
    ("Spoof probability > 0.70?", "Block (automatic termination)"),
]
fy0, fh, fgap = 1.1, 0.78, 0.33
for k, (q, out) in enumerate(flow):
    y = fy0 + k * (fh + fgap)
    novel = k in (1, 2)
    d = label_box(s, 0.6, y, 5.6, fh, q, None, novel=novel, size=13)
    col = GREEN if "Pass" in out else (RED if "Block" in out else AMBER)
    b = box(s, 7.1, y + 0.08, 4.4, fh - 0.16, fill=WHITE, line=col, lw=1.75)
    runs_into(b.text_frame.paragraphs[0], out, 13, col, bold=True)
    b.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    arrow(s, 6.2, y + fh / 2, 7.1, y + fh / 2)
    textbox(s, 6.35, y + fh / 2 - 0.32, 0.6, 0.3, [[("yes", "g")]], size=10)
    if k < len(flow) - 1:
        arrow(s, 3.4, y + fh, 3.4, y + fh + fgap)
        textbox(s, 3.5, y + fh, 0.5, 0.3, [[("no", "g")]], size=10)
yl = fy0 + 5 * (fh + fgap) - fgap
arrow(s, 3.4, yl, 3.4, yl + 0.3)
b = box(s, 0.6, yl + 0.3, 5.6, 0.5, fill=WHITE, line=AMBER, lw=1.75)
runs_into(b.text_frame.paragraphs[0], "Otherwise (0.30–0.70): quality-control review", 13, AMBER, bold=True)
b.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
textbox(s, 11.7, 1.15, 1.45, 5.6, [[("Thresholds as deployed in the demo (round-2 model). The uncertainty threshold was set on validation data for ≤15% genuine review.", "g")]], size=10)

# ============================================================================= 19-20. Module descriptions
s = new_slide("6. Each Module Description: Supporting Front End (Stages 1–5)")
table(s, 0.3, 0.95, 12.7, [2.6, 6.3, 3.8], [
    ["Stage", "Implementation", "Purpose"],
    ["1 Audio input & windowing", "Decode any format to mono 16 kHz; 3-s windows (zero-padded if shorter); long clips scanned in up to 12 windows over the first 60 s", "One standard input; long recordings fully covered"],
    ["2 Robustness augmentation", "Training only: random MP3 (32/64/128 kbps), Gaussian noise at 0/10/20 dB SNR, synthetic room response, 8 kHz band-limiting, or none", "Survive compression, noise, reverberation and telephone audio"],
    ["3 Dual-channel features", "Log-power Mel spectrogram (64 bands, n_fft 1024, hop 512) and 64 MFCCs stacked → 2 × 64 × 94", "Artefact view + vocal-tract view in one input"],
    ["4 Backbone", "Conv 2→32; residual blocks 32→32, 32→64 (stride 2) + SE attention, 64→128 (stride 2); 341,254 parameters in total", "Proven, compact feature learner; CPU real-time"],
    ["5 Shared embedding", "Global average pool → 128-d vector feeding all heads and the novelty scorer", "One representation, consistent heads"],
], size=12, row_h=0.95)
textbox(s, 0.35, 6.55, 12.6, 0.7, [[("These stages use established techniques (slide 13). They are the foundation the novel stages operate on.", "g")]], size=12)

s = new_slide("6. Each Module Description: Novel Stages (6–9)")
legend(s)
table(s, 0.3, 0.95, 12.7, [2.6, 6.3, 3.8], [
    ["Stage", "Implementation", "Why it matters"],
    [[("6 Four parallel heads", "r")], "Spoof head (MLP, BCE loss); attack-type head (5 classes, cross-entropy, weight 0.4, trained jointly); uncertainty = std of P(fake) over N passes with dropout p = 0.35 on the embedding (N = 20 in evaluation, 10 in the demo); novelty = Mahalanobis distance to a Gaussian of genuine training embeddings", "Several independent views of the same evidence"],
    [[("7 Risk-aware gateway", "r")], "Order: minimum duration → novelty → uncertainty → probability bands (0.30 / 0.70). Uncertainty threshold calibrated on validation for ≤15% genuine review (0.445 for the round-2 model)", "Doubtful cases go to people; workload is a parameter"],
    [[("8 Explanation layer", "r")], "Deterministic sentence from the head outputs (live). Per-channel Grad-CAM disagreement map (code only)", "Defensible, reproducible reasons"],
    ["9 Continual learning", "Reviewer-confirmed and high-novelty samples queued for periodic fine-tuning (not built)", "Keeps pace with new generators"],
], size=12, row_h=1.15)

# ============================================================================= 21. Evaluation protocol
s = new_slide("Evaluation Protocol")
slides["Evaluation"] = s
table(s, 0.3, 0.95, 7.3, [2.3, 5.0], [
    ["Item", "Detail"],
    ["Data", "11 public datasets; 40,000-clip pool (20,000 genuine / 20,000 fake) + 4,000 external clips"],
    ["Generators", "40+ systems: WaveFake vocoders, Fake-or-Real TTS, 8 modern TTS (OpenAI, XTTS, VALL-E…), MLAAD, 2025–26 commercial (ElevenLabs, Gemini, Cartesia…)"],
    ["Genuine speech", "Fake-or-Real, LJ Speech, LibriTTS, VCTK (110 speakers), AccentDB (Indian-accent English), Common Voice"],
    ["Split", "Grouped: no recording, sentence or speaker in two splits; 27,363 train / 5,989 val / 6,648 test"],
    ["Unseen generators", "Leave-one-family-out: 9 families held out of training and validation entirely"],
    ["External test", "In-the-Wild: 2,000 genuine + 2,000 deepfakes of public figures, never trained on"],
    ["Robustness", "1,000 test clips re-tested under 8 degradations"],
], size=12, row_h=0.6)
textbox(s, 7.85, 0.95, 5.2, 6.2, [
    ("##", "Why this protocol"),
    [("The students’ original 97.73% came from a random split in which ", ""), ("28%", "b"), (" of validation clips had a copy of the same recording or sentence in training. Under the grouped split the same model design scores ", ""), ("93.2%", "b"), (".", "")],
    "Leave-one-family-out simulates a generator released after training, the threat the gateway is designed for.",
    "In-the-Wild measures performance on real internet audio from a different source than any training data.",
    ("##", "Compute"),
    "Kaggle GPU (T4); full run about 1.8 h. Notebooks and results are in the deliverables folder.",
], size=13, space=6)

# ============================================================================= 22. Technical effect: chart
s = new_slide("Technical Effect: Gateway vs Detector Alone vs Simple Review Rule")
slides["Technical Effect"] = s
lm = matched["logo_mean"]
cats = ["Never-seen generators (mean of 9)", "In-the-Wild (external)", "Clean test"]
det = [summary["gateway_unseen_mean"][0]["fakes_auto_passed"], summary["gateway_in_the_wild"][0]["fakes_auto_passed"], summary["gateway_main_test"][0]["fakes_auto_passed"]]
simple = [lm["base_fakes_passed"] / 100, matched["sets"]["in_the_wild"]["score_margin_same_review"]["fakes_auto_passed"] / 100, matched["sets"]["clean_test"]["score_margin_same_review"]["fakes_auto_passed"] / 100]
gwv = [lm["gw_fakes_passed"] / 100, matched["sets"]["in_the_wild"]["gateway"]["fakes_auto_passed"] / 100, matched["sets"]["clean_test"]["gateway"]["fakes_auto_passed"] / 100]
det_b = [summary["gateway_unseen_mean"][0]["genuine_auto_blocked"], summary["gateway_in_the_wild"][0]["genuine_auto_blocked"], summary["gateway_main_test"][0]["genuine_auto_blocked"]]
simple_b = [lm["base_gen_blocked"] / 100, matched["sets"]["in_the_wild"]["score_margin_same_review"]["genuine_auto_blocked"] / 100, matched["sets"]["clean_test"]["score_margin_same_review"]["genuine_auto_blocked"] / 100]
gw_b = [lm["gw_gen_blocked"] / 100, matched["sets"]["in_the_wild"]["gateway"]["genuine_auto_blocked"] / 100, matched["sets"]["clean_test"]["gateway"]["genuine_auto_blocked"] / 100]


def add_chart(s, x, title, series):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, [round(100 * v, 1) for v in vals])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(0.95), Inches(6.25), Inches(4.4), cd)
    ch = gf.chart
    ch.has_title = True
    ch.chart_title.text_frame.text = title
    tp = ch.chart_title.text_frame.paragraphs[0]
    tp.runs[0].font.size, tp.runs[0].font.bold, tp.runs[0].font.name = Pt(14), True, FONT
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size, ch.legend.font.name = Pt(11), FONT
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    ch.value_axis.maximum_scale = 40
    ch.category_axis.tick_labels.font.size = Pt(10)
    ch.category_axis.tick_labels.font.name = FONT
    for ser, col in zip(ch.plots[0].series, [RGBColor(0xA6, 0xA6, 0xA6), BLUE_LINE, RED]):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = col
    pl = ch.plots[0]
    pl.has_data_labels = True
    dl = pl.data_labels
    dl.number_format, dl.number_format_is_linked = '0.0"%"', False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size, dl.font.name = Pt(10), FONT


add_chart(s, 0.3, "Fakes automatically passed (%)", [("Detector alone (threshold 0.5)", det), ("Simple review rule, same workload", simple), ("DeepShield gateway (current rules)", gwv)])
add_chart(s, 6.75, "Genuine automatically blocked (%)", [("Detector alone (threshold 0.5)", det_b), ("Simple review rule, same workload", simple_b), ("DeepShield gateway (current rules)", gw_b)])
textbox(s, 0.35, 5.5, 12.6, 1.8, [
    [("Reading: ", "b"), ("adding a review lane at ~14% genuine workload (27.5% on In-the-Wild) roughly halves both error types versus the detector alone. ", ""),
     ("But a simple rule that reviews the clips whose spoof score is closest to 0.5, at the same workload, does as well or slightly better than the current gateway rules. ", "r"),
     ("The measured benefit therefore comes mainly from the review lane itself, not yet from the extra heads.", "")],
    [("Simple review rule = review the clips with the smallest |logit P(fake)| until the genuine review rate matches the gateway’s.", "g")],
], size=13, space=6)

# ============================================================================= 23. Specialist contribution
s = new_slide("Technical Effect: Do the Extra Heads Add Information?")
fam_names = {"conformer_fastspeech2": "Conformer-FastSpeech2", "for_tts": "Fake-or-Real TTS", "hifigan": "HiFi-GAN", "melgan_family": "MelGAN family", "mlaad_tts": "MLAAD open TTS", "modern_tts": "Modern TTS (OpenAI, XTTS…)", "parallel_wavegan": "Parallel WaveGAN", "tts_2026": "2025–26 commercial TTS", "waveglow": "WaveGlow"}
rows = [["Held-out (never-seen) family", "Score margin only", "Score + uncertainty + novelty", "Genuine auto-blocked: margin → combined"]]
for r in combined["logo"]:
    better = r["auc_combined"] > r["auc_margin"]
    rows.append([fam_names.get(r["family"], r["family"]), f"{r['auc_margin']:.3f}", [(f"{r['auc_combined']:.3f}", "b" if better else "")], f"{r['margin_gen_blocked']:.1f}% → {r['combined_gen_blocked']:.1f}%"])
mean = combined["mean"]
rows.append([[("Mean", "b")], [(f"{mean['auc_margin']:.3f}", "b")], [(f"{mean['auc_combined']:.3f}", "rb")], [(f"{mean['margin_gen_blocked']:.1f}% → {mean['combined_gen_blocked']:.1f}%", "b")]])
table(s, 0.3, 0.95, 8.2, [2.9, 1.6, 1.9, 1.8], rows, size=11, row_h=0.42)
textbox(s, 8.75, 0.95, 4.3, 6.3, [
    ("##", "What we measured"),
    "How well each signal spots the detector’s wrong verdicts (AUC; 0.5 = chance, 1.0 = perfect).",
    "The combination’s weights were fitted only on the familiar-generator test set, then applied unchanged to each never-seen family. No unseen data was used to fit it.",
    ("##", "Finding"),
    [("The heads ", ""), ("do", "b"), (" carry extra information on never-seen generators: better in ", ""), ("9 of 9", "rb"), (" families, but the gain is small (mean 0.79 → 0.81). Fakes passed at equal workload: 17.0% → 17.4% (no gain); genuine auto-blocked 3.8% → 3.4%.", "")],
    "On familiar generators the heads add nothing over the score margin (0.853 vs 0.850).",
], size=12, space=6)

# ============================================================================= 24. Live inference examples
s = new_slide("Sample Results: Live Inference (round-2 model, laptop CPU)")
slides["Sample Results"] = s
cards = [
    ("result_genuine_pass_card.png", "Genuine (Mark Zuckerberg interview) → Stage 1 Pass", GREEN),
    ("result_deepfake_block_card.png", "AI-generated (Barack Obama) → Stage 2 Block", RED),
    ("result_review_uncertain_card.png", "AI-generated (Zuckerberg) → Stage 3 review, high uncertainty", AMBER),
]
for k, (f, cap, col) in enumerate(cards):
    x = 0.35 + k * 4.25
    picture(s, os.path.join(SCREENS, f), x + 0.25, 0.95, h=4.85)
    textbox(s, x, 5.9, 4.0, 0.7, [[(cap, "b")]], size=12, color=col, align=PP_ALIGN.CENTER)
textbox(s, 0.35, 6.65, 12.6, 0.7, [[("Clips from the In-the-Wild dataset (never used in training). Each card shows the verdict, risk level, P(fake), uncertainty, attack-type distribution, novelty vs threshold, gateway route and the rationale.", "g")]], size=11)

# ============================================================================= 25. Results summary table
s = new_slide("Results Summary (round-2 model)")
rows = [["Held-out family", "Seen EER", "Unseen EER", "Unseen fakes passed", "Flagged novel"]]
for fam in ["tts_2026", "modern_tts", "mlaad_tts", "for_tts", "conformer_fastspeech2", "melgan_family", "hifigan", "waveglow", "parallel_wavegan"]:
    r = LOGO[fam]
    rows.append([fam_names[fam], pct(r["seen_eer"]), pct(r["unseen_eer"]), pct(r["unseen_false_negative_rate"]), pct(r["unseen_fakes_flagged_novel"])])
table(s, 0.3, 0.95, 7.6, [2.9, 1.1, 1.2, 1.3, 1.1], rows, size=11, row_h=0.42)
rb = {r["condition"]: r for r in []}
table(s, 8.15, 0.95, 4.85, [2.9, 1.95], [
    ["Headline metric", "Value"],
    ["Clean test accuracy (6,648 clips)", f"{100 * M['accuracy']:.1f}%"],
    ["Clean test ROC-AUC / EER", f"{M['auc']:.3f} / {100 * M['eer']:.1f}%"],
    ["Genuine flagged fake / fakes passed", f"{100 * M['false_positive_rate']:.1f}% / {100 * M['false_negative_rate']:.1f}%"],
    ["In-the-Wild accuracy / EER", f"{100 * ITW['accuracy']:.1f}% / {100 * ITW['eer']:.1f}%"],
    ["Attack-type accuracy (test fakes)", f"{100 * summary['attack_attribution_accuracy']:.1f}%"],
    ["Uncertainty AUC for spotting errors", f"{summary['uncertainty']['auc_uncertainty_flags_errors']:.2f}"],
    ["Random-split accuracy (for comparison)", "95.8% (28% leakage)"],
], size=11, row_h=0.42)
textbox(s, 8.15, 4.6, 4.85, 2.4, [
    ("##", "Robustness (EER, 1,000 test clips)"),
    "Clean 7.4% · MP3 32 kbps 9.0% · MP3 128 kbps 7.5% · Telephone 8 kHz 10.6% · Reverb 11.8% · Noise 20 dB 13.6% · 10 dB 19.1% · 0 dB 21.2%",
], size=12, space=4)
textbox(s, 0.35, 5.35, 7.6, 1.9, [
    [("EER = equal error rate (lower is better). ", "g"), ("Seen", "g"), (" = the other families on the same genuine test audio. Strong generalisation to unseen vocoders; ", "g"), ("unseen modern and Fake-or-Real TTS remain the weakest cases", "r"), (" (most of their fakes pass when the family is never seen).", "g")],
], size=11)

# ============================================================================= 26. Limitations
s = new_slide("Limitations, Not Yet Reduced to Practice, and Next Experiments")
slides["Limitations"] = s
textbox(s, 0.35, 0.95, 6.3, 6.3, [
    ("##", "Limitations (measured)"),
    ("-", "Never-seen modern TTS: about 2 in 3 fakes pass when the family was never seen"),
    ("-", "Heavy noise: EER rises from 7% (clean) to 21% at 0 dB SNR"),
    ("-", "Clips under 2 s are unreliable; routed to review by design"),
    ("-", "LibriTTS genuine speech: 26% false alarms (other genuine sources 0.4–8%)"),
    ("-", "Manipulated genuine audio (splicing, voice conversion) not tested; the splice/edit class has no training data"),
    ("##", "Not yet reduced to practice"),
    ("-", "Continual-learning loop (claim 13)"),
    ("-", "Per-channel disagreement explanation (claim 11): code only"),
    ("-", "Combined risk-score routing (claim 6): offline analysis only"),
], size=13, space=5)
textbox(s, 6.85, 0.95, 6.2, 6.3, [
    ("##", "Experiments that would strengthen the filing"),
    ("-", [("Replace the hand-written rules with the fitted combined risk score (claim 6), calibrated on validation; re-run leave-one-family-out. ", ""), ("No retraining, no architecture change.", "b")]),
    ("-", "Report error rates at equal review workload against the score-only rule as the standard baseline"),
    ("-", "Evaluate the per-channel disagreement map with reviewers, or drop claim 11"),
    ("-", "Add a voice-conversion / splicing test set to support the attack-type claim"),
    ("-", "Formal prior-art search, especially selective classification combined with OOD detection for anti-spoofing"),
    ("##", "Bottom line for the patent team"),
    [("The integrated system is built and works; ", ""), ("the inventive step currently rests on a modest measured effect", "r"), (". Filing is defensible as a combination claim, but the experiments above should come first if time allows.", "")],
], size=13, space=5)

# ============================================================================= TOC
toc_items = [
    ("Recommendation Summary", "Recommendation Summary"),
    ("Opportunity Review (significance, current solutions, proposed engine, why it works)", "Opportunity Review"),
    ("Independent and Dependent Claims", "Independent Claims"),
    ("Main Patentable Points (ranked by evidence)", "Main Patentable Points"),
    ("Novelty Map", "Novelty Map"),
    ("Prior Art Comparison and Difference from Closest Prior Art", "Prior Art Comparison"),
    ("System Architecture & Module Description", "System Architecture"),
    ("Evaluation Protocol", "Evaluation"),
    ("Technical Effect (evidence)", "Technical Effect"),
    ("Sample Results and Results Summary", "Sample Results"),
    ("Limitations and Next Experiments", "Limitations"),
]
# move 'thanks' to the end
sldIdLst.remove(kept["thanks"])
sldIdLst.append(kept["thanks"])
order = [prs.slides.index(sl) + 1 for sl in prs.slides]  # noqa: F841 (index check)
num = {name: prs.slides.index(sl) + 1 for name, sl in slides.items()}
rows = [["Title", "Slide No."]] + [[t, f"Slide {num[k]}"] for t, k in toc_items]
table(toc, 0.5, 1.0, 12.3, [10.3, 2.0], rows, size=14, row_h=0.45)

prs.core_properties.title = "DeepShield IP Pre-DOI v1.0"
prs.save(OUT)
print("wrote", OUT, "slides:", len(prs.slides))
